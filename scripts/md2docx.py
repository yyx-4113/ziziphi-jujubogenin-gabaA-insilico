#!/usr/bin/env python
"""Convert the manuscript markdown to a submission-ready .docx.

Handles: headings, bold/italic/code, <sub>/<sup>, pipe tables, bullet and
numbered lists, and embeds the three figures with their captions in place of
the "Figure legends" section.

Usage:
    python md2docx.py <input.md> <output.docx> [figdir]
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(11)


# ---------------------------------------------------------------- inline runs
def _emit_text(par, text):
    """Emit plain text, converting any residual HTML entities."""
    text = text.replace("&nbsp;", " ").replace("&amp;", "&")
    run = par.add_run(text)
    return run


def add_inline(par, text):
    """Parse a limited inline markdown subset into runs on `par`."""
    token = re.compile(
        r"(?P<bold>\*\*.+?\*\*)"
        r"|(?P<italic>\*(?!\*).+?\*)"
        r"|(?P<code>`[^`]+`)"
        r"|(?P<sub><sub>.+?</sub>)"
        r"|(?P<sup><sup>.+?</sup>)",
        re.DOTALL,
    )
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            _emit_text(par, text[pos : m.start()])
        if m.group("bold"):
            run = par.add_run(m.group("bold")[2:-2])
            run.bold = True
        elif m.group("italic"):
            run = par.add_run(m.group("italic")[1:-1])
            run.italic = True
        elif m.group("code"):
            run = par.add_run(m.group("code")[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif m.group("sub"):
            run = par.add_run(m.group("sub")[5:-6])
            run.font.subscript = True
        elif m.group("sup"):
            run = par.add_run(m.group("sup")[5:-6])
            run.font.superscript = True
        pos = m.end()
    if pos < len(text):
        _emit_text(par, text[pos:])


# ------------------------------------------------------------------- tables
def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, cell_text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_inline(p, cell_text)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(10)
    for row in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            if i >= len(cells):
                continue
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline(p, cell_text)
            for r in p.runs:
                r.font.size = Pt(10)
    doc.add_paragraph()
    return table


# --------------------------------------------------------------- figures
FIGURES = [
    ("Figure 1", "Fig1_poses.png", 15.0),
    ("Figure 2", "Fig2_md_stability.png", 13.0),
    ("Figure 3", "Fig3_alanine_scan.png", 13.0),
]


def add_figures(doc, figdir, captions):
    doc.add_heading("Figures", level=1)
    for label, fname, width_cm in FIGURES:
        path = os.path.join(figdir, fname)
        if not os.path.exists(path):
            p = doc.add_paragraph()
            p.add_run("[%s: %s not found]" % (label, fname))
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Cm(width_cm))
        cap = captions.get(label)
        if cap:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(cp, cap)
            for r in cp.runs:
                r.font.size = Pt(10)


def parse_figure_legends(lines):
    """Pull '**Figure N.** caption' blocks out of the legend section."""
    captions = {}
    current = None
    for line in lines:
        m = re.match(r"\*\*(Figure \d+)\.\*\*\s*(.*)", line.strip())
        if m:
            current = m.group(1)
            captions[current] = m.group(1) + ". " + m.group(2)
        elif current and line.strip():
            captions[current] += " " + line.strip()
    return captions


# ------------------------------------------------------------------- main
def convert(md_path, out_path, figdir):
    with open(md_path, "r", encoding="utf-8") as fh:
        raw_lines = fh.read().split("\n")

    # Pre-scan figure legends so captions can travel with the images.
    legends, in_legend = {}, False
    for line in raw_lines:
        if line.startswith("## Figure legends"):
            in_legend = True
            continue
        if in_legend and line.startswith("## "):
            in_legend = False
        if in_legend:
            for k, v in parse_figure_legends([line]).items():
                legends[k] = (legends.get(k, "") + " " + v).strip()

    doc = Document()

    # Page + base style setup
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for lvl, size in ((1, 15), (2, 13), (3, 12)):
        st = doc.styles["Heading %d" % lvl]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = True
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    n = len(raw_lines)
    while i < n:
        line = raw_lines[i]
        stripped = line.strip()

        # skip blank lines and horizontal rules
        if not stripped or stripped == "---":
            i += 1
            continue

        # figure legends section -> render real figures with captions
        if stripped.startswith("## Figure legends"):
            add_figures(doc, figdir, legends)
            i += 1
            while i < n and not raw_lines[i].startswith("## "):
                i += 1
            continue

        # table block
        if stripped.startswith("|") and i + 1 < n and re.match(
            r"^\|[\s:\-\|]+\|$", raw_lines[i + 1].strip()
        ):
            rows = [split_row(stripped)]
            i += 2
            while i < n and raw_lines[i].strip().startswith("|"):
                rows.append(split_row(raw_lines[i].strip()))
                i += 1
            add_table(doc, rows)
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                # Title: centred, larger
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text)
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(15)
                p.paragraph_format.space_after = Pt(12)
            else:
                h = doc.add_heading("", level=min(level - 1, 3))
                add_inline(h, text)
                for r in h.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # bullet list
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            i += 1
            continue

        # numbered list
        m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1
            continue

        # ordinary paragraph (merge soft-wrapped continuation lines)
        buf = [stripped]
        i += 1
        while i < n:
            nxt = raw_lines[i].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("- ")
                or re.match(r"^\d+\.\s", nxt)
            ):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(buf))

    doc.save(out_path)
    print("wrote %s" % out_path)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "manuscript.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "manuscript.docx"
    figs = sys.argv[3] if len(sys.argv) > 3 else "figures"
    convert(src, dst, figs)
